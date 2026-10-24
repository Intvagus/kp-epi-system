"""Tests for the "Divisional Officers' Consolidated Monthly Summary" workbook
(src/pipeline/divisional_officer_reports.py) -- a genuinely new report type
found from a real user upload (2026-10), "Divisional_Officers_Summary_
Report_September_2026.xlsx": one narrative (not Yes/No/N-A) monthly report
sheet per officer (DO-1..DO-N), entirely different from the existing Admin
Activities compliance checklist even though it shares several real officer
names. Per the user's explicit choice, folded into the dashboard's existing
Divisional Officer Compliance section, not a new tab.

Uses small synthetic workbooks built inline (not the real uploaded file,
which contains real narrative field reports and was never committed to
this repo), so these tests run unconditionally.
"""
import io

import openpyxl
import pytest

from src.pipeline.detect import detect_workbook_type
from src.pipeline.divisional_officer_reports import (
    build_summary, find_divisional_officer_reports_files, load_divisional_officer_reports,
)


def _status_sheet(wb):
    ws = wb.create_sheet("Status")
    ws.cell(row=1, column=1, value="Divisional Officers' Reports - Status & How to Use")
    return ws


def _officer_sheet(wb, sheet_name, name, month, duty_station, districts, sections):
    ws = wb.create_sheet(sheet_name)
    ws.cell(row=1, column=1, value=name)
    ws.cell(row=1, column=2, value="Reporting Month:")
    ws.cell(row=1, column=3, value=month)
    ws.cell(row=2, column=2, value="Duty Station:")
    ws.cell(row=2, column=3, value=duty_station)
    ws.cell(row=3, column=2, value="Assigned Districts:")
    ws.cell(row=3, column=3, value=districts)
    r = 4
    headings = {
        "key_highlights": "Key Highlights of the assigned Districts: (For example, EPI Performance)",
        "key_activities": "Brief Description of Key activities/field visits performed, meetings attended.",
        "key_challenges": " Key Challenges (District wise brief description, Way forward if any)",
        "initiatives": "Initiative taken at District level (if any).",
        "other_business": "Any other Business: (Innovations, suggestions).",
    }
    for key, heading in headings.items():
        items = sections.get(key, [])
        ws.cell(row=r, column=2, value=heading)
        if items:
            ws.cell(row=r, column=3, value=items[0])
            r += 1
            for item in items[1:]:
                ws.cell(row=r, column=3, value=item)
                r += 1
        else:
            r += 1
    return ws


def _build_workbook(officers) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _status_sheet(wb)
    for i, o in enumerate(officers, 1):
        _officer_sheet(wb, f"DO-{i}", o["name"], o.get("month"), o.get("duty_station"),
                        o.get("districts"), o.get("sections", {}))
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def _save(tmp_path, name, content: bytes):
    path = tmp_path / name
    path.write_bytes(content)
    return path


def test_detected_as_divisional_officer_reports_not_admin_activities(tmp_path):
    content = _build_workbook([
        {"name": "Dr A", "month": "September 2026", "duty_station": "X", "districts": "Y",
         "sections": {"key_highlights": ["• did a thing"]}},
    ])
    path = _save(tmp_path, "report.xlsx", content)
    result = detect_workbook_type(path)
    assert result.workbook_type == "divisional_officer_reports"


def test_find_files_locates_real_workbook(tmp_path):
    content = _build_workbook([{"name": "Dr A"}])
    _save(tmp_path, "report.xlsx", content)
    files = find_divisional_officer_reports_files(tmp_path)
    assert len(files) == 1


def test_find_files_skips_unrelated_workbook(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.cell(row=1, column=1, value="Not a divisional officer report")
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    _save(tmp_path, "other.xlsx", buf.read())
    assert find_divisional_officer_reports_files(tmp_path) == []


def test_loads_sections_and_meta_correctly(tmp_path):
    content = _build_workbook([
        {"name": "Dr Imran Khan", "month": "September 2026", "duty_station": "Hangu",
         "districts": "Kohat, Hangu",
         "sections": {
             "key_highlights": ["• highlight one", "• highlight two"],
             "key_activities": ["• activity one"],
         }},
    ])
    path = _save(tmp_path, "report.xlsx", content)
    loaded = load_divisional_officer_reports(path)
    assert len(loaded["officers"]) == 1
    o = loaded["officers"][0]
    assert o["name"] == "Dr Imran Khan"
    assert o["duty_station"] == "Hangu"
    assert o["assigned_districts"] == "Kohat, Hangu"
    # Raw text is kept as-is, including the source's own "• " bullet prefix
    # -- stripped only at display time in template.html (same convention as
    # WHO Activities' narrative sections), never altered in the pipeline.
    assert o["sections"]["key_highlights"] == ["• highlight one", "• highlight two"]
    assert o["sections"]["key_activities"] == ["• activity one"]
    assert o["sections"]["key_challenges"] == []
    assert o["entry_count"] == 3
    assert o["has_any_content"] is True


def test_officer_with_no_content_flagged_not_fabricated(tmp_path):
    """An officer sheet with a name but nothing else (no duty station, no
    bullets) -- the real "Dr Haroon Rashid" (DO-7) case in the file this
    was built from -- must be reported honestly as not-yet-reported, never
    silently dropped or given fabricated content."""
    content = _build_workbook([
        {"name": "Dr Nobody", "duty_station": None, "districts": None, "sections": {}},
    ])
    path = _save(tmp_path, "report.xlsx", content)
    loaded = load_divisional_officer_reports(path)
    o = loaded["officers"][0]
    assert o["name"] == "Dr Nobody"
    assert o["entry_count"] == 0
    assert o["has_any_content"] is False
    summary = build_summary(loaded)
    assert summary["not_reported_count"] == 1
    assert summary["reported_count"] == 0


def test_reporting_month_mismatch_flagged_not_silently_normalized():
    loaded = {"source_file": "x.xlsx", "officers": [
        {"name": "A", "reporting_month_label": "September 2026", "entry_count": 1, "has_any_content": True,
         "sections": {}, "duty_station": None, "assigned_districts": None},
        {"name": "B", "reporting_month_label": "September 2026", "entry_count": 1, "has_any_content": True,
         "sections": {}, "duty_station": None, "assigned_districts": None},
        {"name": "C", "reporting_month_label": "September 2024", "entry_count": 1, "has_any_content": True,
         "sections": {}, "duty_station": None, "assigned_districts": None},
    ]}
    summary = build_summary(loaded)
    assert summary["common_reporting_month"] == "September 2026"
    assert summary["reporting_month_mismatches"] == ["C"]


def test_unmatched_section_heading_does_not_bleed_into_previous_section(tmp_path):
    """A row whose column-B text doesn't match any of the 5 known section
    markers (e.g. a renamed heading in a future export) must end the
    current section rather than silently keep attributing bullets to it."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _status_sheet(wb)
    ws = wb.create_sheet("DO-1")
    ws.cell(row=1, column=1, value="Dr A")
    ws.cell(row=4, column=2, value="Key Highlights of the assigned Districts: (example)")
    ws.cell(row=4, column=3, value="real highlight")
    ws.cell(row=5, column=2, value="Some Unrecognized New Heading")
    ws.cell(row=5, column=3, value="should not count as a highlight")
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    path = _save(tmp_path, "report.xlsx", buf.read())
    loaded = load_divisional_officer_reports(path)
    o = loaded["officers"][0]
    assert o["sections"]["key_highlights"] == ["real highlight"]
    assert o["entry_count"] == 1
    assert o["unmatched_headings"] == ["Some Unrecognized New Heading"]


def test_no_officer_sheets_raises_clear_error(tmp_path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _status_sheet(wb)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    path = _save(tmp_path, "report.xlsx", buf.read())
    with pytest.raises(ValueError, match="No officer sheets"):
        load_divisional_officer_reports(path)
