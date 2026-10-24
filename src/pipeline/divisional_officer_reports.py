"""Read the "Divisional Officers' Consolidated Monthly Summary" workbook --
a per-officer NARRATIVE monthly activity report (Key Highlights, Key
Activities/Field Visits, Key Challenges, Initiatives Taken, Any Other
Business), one sheet per officer (DO-1..DO-N). This is entirely different
from the existing Admin Activities Yes/No/N-A compliance checklist
(admin_activities.py) even though it covers the same Divisional Officer
role and reuses several of the same real officer names (see CLAUDE.md) --
confirmed by direct inspection of the real uploaded file
("Divisional_Officers_Summary_Report_September_2026.xlsx"). Per the user's
explicit choice, this is folded into the dashboard's existing "Divisional
Officer Compliance" section as an additional "Monthly Activity Reports"
card, not a new top-level tab.

Confirmed by direct inspection, nothing guessed:
- A "Status" sheet (A1: "Divisional Officers' Reports - Status & How to
  Use") documents the workbook's own structure and lists every officer
  tab's reporting status in a small roster table -- used here only to
  confirm detection (find_divisional_officer_reports_files), never read
  for real values. Several of its own cells (Duty Stations/Assigned
  Districts/Reports Received/Summary rows used) are themselves literal
  Excel #N/A error values in the file as received (broken/uncalculated
  formulas) -- a real, visible issue in the source file, not something
  this module works around by reading that sheet anyway.
- One "DO-N" sheet per officer, each identically shaped: officer name in
  A1, Reporting Month in C1, Duty Station in C2, Assigned Districts in C3,
  then 5 fixed section headings in column B (long, instructional text --
  matched by a short, recognizable substring, not the full text, since the
  real headings carry parenthetical examples that could drift slightly
  between months), each followed by zero or more bullet entries in column
  C. Real per-officer data is read directly from these DO-N sheets, never
  from the Status/Summary sheets' formula-derived values.
- A "Summary" sheet auto-aggregates every officer's entries via Excel
  formulas (tagged by officer/duty station) -- not read by this module,
  for the same reason: it's a live formula view of the DO-N sheets, which
  this module already reads directly and independently.
- An "Engine" sheet (1,759 rows in the sample received) backs the Summary
  sheet's formulas -- not used.
- A real data-quality issue found in the sample received: one officer
  ("Dr Mansoor") showed a different Reporting Month than every other
  officer tab (a likely source-side typo, e.g. year digit) -- surfaced
  here as `reporting_month_mismatches`, not silently normalized away.
"""
import re
from collections import Counter
from pathlib import Path

import openpyxl

from .sheet_utils import resolve_sheet_name

STATUS_SHEET = "Status"
STATUS_TITLE_MARKER = "divisional officers' reports - status"

_DO_SHEET_RE = re.compile(r"^do-?\s*(\d+)$", re.I)

# Matched as a substring of the real (long, instructional) column-B heading
# text, lowercased -- not an exact match, since the real headings carry
# parenthetical examples that could drift slightly between monthly exports.
SECTION_HEADING_MARKERS = [
    ("key_highlights", "Key Highlights of the Assigned Districts", "key highlights of the assigned districts"),
    ("key_activities", "Key Activities / Field Visits", "brief description of key activities"),
    ("key_challenges", "Key Challenges", "key challenges"),
    ("initiatives", "Initiatives Taken", "initiative taken at district level"),
    ("other_business", "Any Other Business", "any other business"),
]
SECTION_KEYS = [key for key, _, _ in SECTION_HEADING_MARKERS]
SECTION_LABELS = {key: label for key, label, _ in SECTION_HEADING_MARKERS}


def find_divisional_officer_reports_files(raw_dir: Path) -> list[Path]:
    files = []
    for path in sorted(raw_dir.glob("*.xlsx")):
        try:
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        except Exception:
            continue
        status_name = resolve_sheet_name(wb.sheetnames, STATUS_SHEET)
        if status_name is None:
            continue
        title = wb[status_name].cell(row=1, column=1).value
        if title and STATUS_TITLE_MARKER in str(title).strip().lower():
            files.append(path)
    return files


def _reporting_month_label(value) -> str | None:
    if value is None:
        return None
    try:
        return value.strftime("%B %Y")
    except AttributeError:
        return str(value).strip() or None


def _load_officer_sheet(ws, sheet_name: str) -> dict:
    name = ws.cell(row=1, column=1).value
    name = str(name).strip() if name else None

    sections = {key: [] for key in SECTION_KEYS}
    unmatched_headings = []
    current_key = None
    max_row = ws.max_row or 0
    for r in range(4, max_row + 1):
        heading_cell = ws.cell(row=r, column=2).value
        entry_cell = ws.cell(row=r, column=3).value
        if heading_cell:
            heading_norm = str(heading_cell).strip().lower()
            matched = next((key for key, _, marker in SECTION_HEADING_MARKERS if marker in heading_norm), None)
            if matched is None:
                unmatched_headings.append(str(heading_cell).strip())
            current_key = matched
        if current_key and entry_cell and str(entry_cell).strip():
            sections[current_key].append(str(entry_cell).strip())

    entry_count = sum(len(v) for v in sections.values())
    duty_station = ws.cell(row=2, column=3).value
    assigned_districts = ws.cell(row=3, column=3).value
    duty_station = str(duty_station).strip() if duty_station else None
    assigned_districts = str(assigned_districts).strip() if assigned_districts else None

    return {
        "sheet": sheet_name,
        "name": name,
        "reporting_month_label": _reporting_month_label(ws.cell(row=1, column=3).value),
        "duty_station": duty_station,
        "assigned_districts": assigned_districts,
        "sections": sections,
        "entry_count": entry_count,
        "unmatched_headings": unmatched_headings,
        "has_any_content": bool(entry_count or duty_station or assigned_districts),
    }


def load_divisional_officer_reports(path: Path) -> dict:
    """Returns {'source_file': ..., 'officers': [...]} -- one entry per
    real (non-blank) DO-N sheet, read directly, never via the Status/
    Summary sheets' own (sometimes #N/A) formula-derived values."""
    wb = openpyxl.load_workbook(path, data_only=True)
    do_sheets = sorted(
        (n for n in wb.sheetnames if _DO_SHEET_RE.match(n.strip())),
        key=lambda n: int(_DO_SHEET_RE.match(n.strip()).group(1)),
    )
    if not do_sheets:
        raise ValueError(
            f"No officer sheets (expected names like 'DO-1', 'DO-2', ...) found in {path.name}."
        )

    officers = []
    for sheet_name in do_sheets:
        officer = _load_officer_sheet(wb[sheet_name], sheet_name)
        if officer["name"]:
            officers.append(officer)

    return {"source_file": path.name, "officers": officers}


def build_summary(loaded: dict) -> dict:
    officers = loaded["officers"]
    reported = [o for o in officers if o["has_any_content"]]
    months = [o["reporting_month_label"] for o in officers if o["reporting_month_label"]]
    common_month = Counter(months).most_common(1)[0][0] if months else None
    mismatches = [
        o["name"] for o in officers
        if o["reporting_month_label"] and o["reporting_month_label"] != common_month
    ]

    return {
        "status": "ok",
        "source_file": loaded["source_file"],
        "officer_count": len(officers),
        "reported_count": len(reported),
        "not_reported_count": len(officers) - len(reported),
        "total_entries": sum(o["entry_count"] for o in officers),
        "common_reporting_month": common_month,
        "reporting_month_mismatches": mismatches,
        "officers": officers,
    }


PROJECT_ROOT = Path(__file__).resolve().parents[2]
PROCESSED_DIR = PROJECT_ROOT / "data" / "processed"


def run_divisional_officer_reports(raw_dir: Path | None = None, processed_dir: Path | None = None) -> dict | None:
    """Independent of every other domain, including the Admin Activities
    checklist this folds in alongside on the dashboard -- see every other
    run_* entry point in this project for the same convention. Returns None
    (writes nothing) if no matching workbook is found."""
    import json

    raw_dir = raw_dir or (PROJECT_ROOT / "data" / "raw")
    processed_dir = processed_dir or PROCESSED_DIR
    files = find_divisional_officer_reports_files(raw_dir)
    if not files:
        return None

    print("\nDivisional Officer monthly reports pipeline starting...")
    path = files[0]
    print(f"  Loading {path.name}...")
    try:
        loaded = load_divisional_officer_reports(path)
    except ValueError as e:
        raise SystemExit(f"\nFAILED loading {path.name}: {e}") from e
    summary = build_summary(loaded)
    print(f"  {summary['officer_count']} officer sheets found, {summary['reported_count']} with content, "
          f"{summary['total_entries']} total entries")

    processed_dir.mkdir(parents=True, exist_ok=True)
    with open(processed_dir / "divisional_officer_reports_summary.json", "w", encoding="utf-8") as f:
        json.dump(summary, f, indent=2, default=str)
    print("Divisional Officer monthly reports pipeline finished OK.")
    return summary
